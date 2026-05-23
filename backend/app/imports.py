from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select

from .auth import get_password_hash, requires_role
from .db import get_session
from .models import Role, School, Student, User

router = APIRouter()


REQUIRED_HEADERS = ["username", "full_name", "email"]
OPTIONAL_HEADERS = ["password", "tenant_id"]
ALL_HEADERS = REQUIRED_HEADERS + OPTIONAL_HEADERS


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _to_int(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(value)
    except ValueError:
        return None


def _extract_rows(file_bytes: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        raise HTTPException(status_code=400, detail="Excel file is empty")

    headers = [_normalize_header(h) for h in raw_rows[0]]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise HTTPException(status_code=400, detail=f"Missing required headers: {', '.join(missing)}")

    unknown = [h for h in headers if h and h not in ALL_HEADERS]
    if unknown:
        raise HTTPException(status_code=400, detail=f"Unknown headers: {', '.join(unknown)}")

    header_index = {h: i for i, h in enumerate(headers) if h}
    valid_rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for row_number, values in enumerate(raw_rows[1:], start=2):
        row = {}
        for header, idx in header_index.items():
            row[header] = values[idx] if idx < len(values) else None

        username = (str(row.get("username") or "")).strip()
        email = (str(row.get("email") or "")).strip()
        full_name = (str(row.get("full_name") or "")).strip()
        password = (str(row.get("password") or "")).strip()
        tenant_id = _to_int(row.get("tenant_id"))

        row_errors = []
        if not username:
            row_errors.append("username is required")
        if not email:
            row_errors.append("email is required")
        if not full_name:
            row_errors.append("full_name is required")
        if tenant_id is not None and tenant_id <= 0:
            row_errors.append("tenant_id must be a positive integer")

        if row_errors:
            errors.append({"row": row_number, "errors": row_errors})
            continue

        valid_rows.append(
            {
                "row": row_number,
                "username": username,
                "email": email,
                "full_name": full_name,
                "password": password,
                "tenant_id": tenant_id,
            }
        )

    return valid_rows, errors


@router.post("/imports/students")
def import_students(
    school_id: int,
    file: UploadFile = File(...),
    current=Depends(requires_role(Role.superadmin, Role.admin, Role.school)),
    session: Session = Depends(get_session),
):
    school = session.get(School, school_id)
    if not school:
        raise HTTPException(status_code=404, detail="School not found")

    content = file.file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    try:
        valid_rows, parse_errors = _extract_rows(content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid Excel file: {exc}") from exc
    if not valid_rows and parse_errors:
        return {
            "status": "failed",
            "filename": file.filename,
            "school_id": school_id,
            "processed": 0,
            "created_users": 0,
            "updated_users": 0,
            "created_students": 0,
            "linked_students": 0,
            "errors": parse_errors,
        }

    created_users = 0
    updated_users = 0
    created_students = 0
    linked_students = 0
    import_errors = list(parse_errors)

    for item in valid_rows:
        username = item["username"]
        try:
            existing = session.exec(select(User).where(User.username == username)).first()
            tenant_id = item["tenant_id"] if item["tenant_id"] is not None else school.tenant_id

            if existing:
                existing.email = item["email"]
                existing.full_name = item["full_name"]
                if existing.role != Role.student:
                    existing.role = Role.student
                if tenant_id is not None:
                    existing.tenant_id = tenant_id
                existing.school_id = school_id
                session.add(existing)
                session.flush()
                user = existing
                updated_users += 1
            else:
                raw_password = item["password"] or f"{username}@123"
                user = User(
                    username=username,
                    email=item["email"],
                    full_name=item["full_name"],
                    role=Role.student,
                    tenant_id=tenant_id,
                    school_id=school_id,
                    hashed_password=get_password_hash(raw_password),
                )
                session.add(user)
                session.flush()
                created_users += 1

            student = session.exec(select(Student).where(Student.user_id == user.id)).first()
            if student:
                student.school_id = school_id
                if tenant_id is not None:
                    student.tenant_id = tenant_id
                session.add(student)
                linked_students += 1
            else:
                student = Student(user_id=user.id, school_id=school_id, tenant_id=tenant_id)
                session.add(student)
                created_students += 1
            session.commit()
        except Exception as exc:
            session.rollback()
            import_errors.append({"row": item["row"], "errors": [str(exc)]})

    return {
        "status": "completed",
        "filename": file.filename,
        "school_id": school_id,
        "processed": len(valid_rows),
        "created_users": created_users,
        "updated_users": updated_users,
        "created_students": created_students,
        "linked_students": linked_students,
        "errors": import_errors,
    }


@router.get("/imports/students/template")
def download_students_template(current=Depends(requires_role(Role.superadmin, Role.admin, Role.school))):
    wb = Workbook()
    ws = wb.active
    ws.title = "students"
    ws.append(["username", "full_name", "email", "password", "tenant_id"])
    ws.append(["john.1", "John Doe", "john@example.com", "john@123", ""])
    ws.append(["mary.2", "Mary Stone", "mary@example.com", "", ""])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=students_import_template.xlsx"},
    )
